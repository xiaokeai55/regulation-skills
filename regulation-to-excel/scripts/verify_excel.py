#!/usr/bin/env python3
"""Read-only checks for a regulation workbook; semantic/source review remains mandatory."""
import argparse
import json
import re
import posixpath
from urllib.parse import unquote
from xml.etree import ElementTree as ET
from zipfile import ZipFile, BadZipFile
from collections import defaultdict
from copy import copy
from pathlib import Path
from openpyxl import load_workbook
from language_provenance import language_headers

ENGLISH = {"en", "english", "英文", "英语"}
ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
MERGES = {"A1:A2", "B1:B2", "C1:G1", "H1:H2", "I1:I2", "J1:J2", "K1:K2", "L1:L2", "M1:M2", "N1:N2"}

def check_package(path):
    """Catch package faults that permissive workbook readers may silently ignore."""
    failures = []
    doc_rel = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    main = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    try:
        with ZipFile(path) as package:
            names = package.namelist()
            if len(names) != len(set(names)):
                failures.append('XLSX包内有重复文件名')
            bad = package.testzip()
            if bad:
                failures.append('XLSX包CRC检查失败：' + bad)
            for required in ['[Content_Types].xml', '_rels/.rels', 'xl/workbook.xml']:
                if required not in names:
                    failures.append('XLSX包缺少必需文件：' + required)
            documents = {}
            for name in names:
                if name.endswith(('.xml', '.rels', '.vml')):
                    try:
                        documents[name] = ET.fromstring(package.read(name))
                    except ET.ParseError as error:
                        failures.append(name + ' XML无效：' + str(error))
            for name, document in documents.items():
                if not name.endswith('.rels'):
                    continue
                owner = '' if name == '_rels/.rels' else posixpath.join(
                    posixpath.dirname(posixpath.dirname(name)), posixpath.basename(name)[:-5])
                ids = set()
                annotations = set()
                note_count = 0
                for relation in document:
                    rid = relation.get('Id')
                    if not rid or rid in ids:
                        failures.append(name + '关系ID缺失或重复：' + str(rid))
                    ids.add(rid)
                    if relation.get('TargetMode') == 'External':
                        continue
                    target = relation.get('Target', '')
                    part = posixpath.normpath(posixpath.join(posixpath.dirname(owner), unquote(target))).lstrip('/')
                    if not target or part not in names:
                        failures.append(name + '内部关系目标不存在：' + target)
                        continue
                    kind = relation.get('Type', '').rsplit('/', 1)[-1]
                    root = documents.get(part)
                    if kind == 'comments':
                        note_count += 1
                        if note_count > 1:
                            failures.append(owner + '关联多个备注文件')
                        if root is not None:
                            authors = root.find(f'{{{main}}}authors')
                            author_count = len(authors) if authors is not None else 0
                            for comment in root.findall(f'{{{main}}}commentList/{{{main}}}comment'):
                                ref = comment.get('ref', '').replace('$', '').upper()
                                if not re.fullmatch(r'[A-Z]+[1-9][0-9]*', ref) or ref in annotations:
                                    failures.append(owner + '备注引用无效或重复：' + ref)
                                annotations.add(ref)
                                author = comment.get('authorId', '')
                                if not author.isdigit() or int(author) >= author_count:
                                    failures.append(part + '备注作者引用无效：' + author)
                if owner in documents:
                    for element in documents[owner].iter():
                        for key, value in element.attrib.items():
                            if key.startswith('{' + doc_rel + '}') and value not in ids:
                                failures.append(owner + '引用不存在的关系ID：' + value)
            content_types = documents.get('[Content_Types].xml')
            if content_types is not None:
                for element in content_types:
                    part = element.get('PartName')
                    if part and unquote(part).lstrip('/') not in names:
                        failures.append('内容类型引用不存在的文件：' + part)
    except (BadZipFile, OSError, RuntimeError, ValueError) as error:
        failures.append('无法完整读取XLSX包：' + str(error))
    return failures

def meaningful(value):
    return isinstance(value, str) and bool(value.strip())

def labels(text):
    return re.findall(r"(?m)^\s*((?:[a-z]|[ivx]+|[0-9]+)[.)]|\([a-z]\)|[①②③④⑤⑥⑦⑧⑨⑩])(?:\s|$)", text)

def display_label(label, source_key, need, name):
    if not isinstance(label, dict):
        need(False, name + "缺少原文标题记录")
        return None
    for key in ['original', 'local', 'english', 'chinese']:
        need(meaningful(label.get(key)), name + "缺少标题语言字段：" + key)
    need(label.get('original') == label.get(source_key), name + '原始标题与实际来源语言不一致')
    values = [label.get(key) for key in ["local", "english", "chinese"]]
    return "\n".join(dict.fromkeys(v for v in values if meaningful(v)))

def fill_rgb(fill):
    for color in [fill.fgColor, fill.bgColor]:
        if color.type == "rgb" and color.rgb != "00000000":
            return color.rgb[-6:].upper()
    return None

def applicability_colors(data):
    return data.get("format", {}).get("applicability_colors", {"是": "#00B050", "否": "#BFBFBF"})

def check_manifest(data):
    failures = []
    def need(ok, message):
        if not ok:
            failures.append(message)
    try:
        language_headers(data)
    except (ValueError, KeyError, TypeError, IndexError, AttributeError) as error:
        need(False, str(error))
    reg = data.get("regulation", {})
    for name in ["title", "user_url", "official_landing_url", "actual_text_url", "version", "checked_date", "scope", "source_language", "source_note"]:
        need(meaningful(reg.get(name)), "来源字段缺失：" + name)
    target = data.get("target_environment", {})
    need(isinstance(target, dict), "目标环境信息无效")
    if isinstance(target, dict):
        for name in ["name", "business_context"]:
            need(meaningful(target.get(name)), "目标环境字段缺失：" + name)
    review = data.get("source_review", {})
    need(isinstance(review, dict), "来源核验记录无效")
    if isinstance(review, dict):
        need(review.get("url_status") == "matched", "URL尚未核验匹配，不能生成完成文件")
        need(isinstance(review.get("newer_version_found"), bool), "缺少新版发现情况")
        if review.get("newer_version_found"):
            need(meaningful(review.get("confirmed_version")) and review.get("confirmed_version") == reg.get("version"), "发现新版后缺少与整理版本一致的用户版本选择")
    colors = applicability_colors(data)
    for value in ["是", "否"]:
        need(bool(re.fullmatch(r"#[0-9A-Fa-f]{6}", str(colors.get(value, "")))), "适用性颜色无效：" + value)
    inventory = data.get("inventory", [])
    units = data.get("units", [])
    ids = [u.get("id") for u in units]
    need(bool(inventory) and all(meaningful(x) for x in inventory), "原文清单为空或含无效编号")
    need(len(inventory) == len(set(inventory)), "原文清单存在重复编号")
    need(len(ids) == len(set(ids)), "整理项存在重复编号")
    need(inventory == ids, "原文清单与输出编号／顺序不一致")
    evidence = defaultdict(list)
    for item in data.get("evidence", []):
        evidence[item.get("id")].append(item)
        need(item.get("id") in ids, "证据引用不存在的条款：" + str(item.get("id")))
        for field in ["file", "locator", "quote", "source_text"]:
            need(meaningful(item.get(field)), "证据字段缺失：" + str(item.get("id")) + "/" + field)
        if meaningful(item.get("quote")) and meaningful(item.get("source_text")):
            need(item["quote"] in item["source_text"], "摘录不在读取的实践原文中：" + str(item.get("id")))
    language = data.get('language_plan', {})
    if not isinstance(language, dict):
        language = {}
    for key in ['local_language', 'local_source_search_note']:
        need(meaningful(language.get(key)), '缺少当地语言核验字段：' + key)
    kind = language.get('local_text_kind')
    need(kind in ['official_original', 'translated_from_english'], '缺少有效的当地语言文本来源类型')
    from_english = kind == 'translated_from_english'
    local_is_english = str(language.get('local_language', '')).lower() in ENGLISH
    source_column = 5 if from_english else 4
    source_key = 'english' if from_english else 'local'
    if from_english:
        need(str(reg.get('source_language', '')).lower() in ENGLISH, '当地语言回译必须以核验过的英文为来源')
        need('非官方译文，译自英文' in str(reg.get('source_note', '')), '来源备注缺少当地语言非官方回译标识')
    else:
        need(reg.get('source_language') == language.get('local_language'), '当地官方原文语言与当地语言不一致')
    for index, unit in enumerate(units, 1):
        uid = str(unit.get("id"))
        row = unit.get("values", [])
        need(meaningful(unit.get("source_locator")), uid + "无原文位置")
        need(meaningful(unit.get("original_text")), uid + "无原文")
        need(isinstance(row, list) and len(row) == 14, uid + "不是14列")
        if not isinstance(row, list) or len(row) != 14:
            continue
        need(row[0] == index, uid + "A序号错误")
        source_labels = unit.get("source_labels", {})
        need(isinstance(source_labels, dict), uid + "原文标题记录无效")
        if isinstance(source_labels, dict):
            for field, column in [("control_domain", 1), ("clause_number", 2)]:
                expected = display_label(source_labels.get(field), source_key, need, uid + "/" + field)
                need(row[column] == expected, uid + "多语言标题不完整或与原文标题记录不同：" + chr(65 + column))
            need("control_subdomain" in source_labels, uid + "未记录原文是否存在控制子域")
            subdomain = source_labels.get("control_subdomain")
            if subdomain is None:
                need(row[3] == "/", uid + "无原文控制子域时D必须为/")
            else:
                expected = display_label(subdomain, source_key, need, uid + "/control_subdomain")
                need(row[3] == expected, uid + "D不是记录的原文真实子标题／译文")
        need(row[source_column] == unit.get("original_text"), uid + "实际来源语言栏与原文清单不同")
        need(meaningful(row[4]), uid + "当地语言内容缺失")
        need(meaningful(row[6]), uid + "中文译文缺失")
        need(meaningful(row[5]), uid + "英语内容缺失")
        if local_is_english:
            need(row[4] == row[5], uid + '当地语言为英语时E和F必须完整填写同一英文内容')
        need(row[7] in ["是", "否"], uid + "H不是是／否")
        for j, value in enumerate(row):
            need(value is None or isinstance(value, (str, int, float)), uid + "含不支持的单元格值")
            need(not isinstance(value, str) or len(value) <= 32767, uid + "单元格超Excel字符上限：" + str(j + 1))
        if row[7] == "否":
            need(all(value is None for value in row[8:13]), uid + "否行I至M不是真正空白")
            need(not evidence.get(unit.get("id")), uid + "否行包含实践证据")
        elif row[7] == "是":
            need(all(meaningful(value) for value in row[8:14]), uid + "是行I至N不完整")
            blocks = evidence.get(unit.get("id"), [])
            need(bool(blocks), uid + "是行没有可追溯证据")
            if blocks and all(meaningful(b.get("quote")) for b in blocks):
                numbered = "\n\n".join(f"{n}. {b['quote']}" for n, b in enumerate(blocks, 1))
                plain = "\n\n".join(b["quote"] for b in blocks)
                need(row[9] in [numbered, plain], uid + "J并非证据摘录的完整组合")
                for block in blocks:
                    need(block.get("file", "") in (row[10] or ""), uid + "K缺少证据文件")
                    need(block.get("locator", "") in (row[10] or ""), uid + "K缺少证据定位")
            need((row[11] == "/") == (row[12] == "/"), uid + "L／M无服务标记不一致")
        for col in [4, 5, 6]:
            if meaningful(row[col]):
                need(labels(unit.get('original_text') or '') == labels(row[col]), uid + "原文与译文子项标记不同：" + chr(65 + col))
                need((unit.get('original_text') or '').count("•") == row[col].count("•"), uid + "原文与译文列表项数量不同：" + chr(65 + col))
    lookup = {u.get("id"): u for u in units}
    for item in data.get("translation_checks", []):
        uid = item.get("id")
        column = item.get("column")
        need(uid in lookup and column in ["E", "F", "G"], "无效关键译文检查")
        if uid not in lookup or column not in ["E", "F", "G"]:
            continue
        text = lookup[uid]["values"][ord(column) - 65] or ""
        for expected in item.get("contains", []):
            need(expected in text, str(uid) + "关键译文表达缺失：" + expected)
    return failures

def verify_workbook(data, path, template_path=None):
    failures = check_manifest(data)
    package_failures = check_package(path)
    failures.extend(package_failures)
    limitations = '仅机械校验；不替代法规、翻译及证据复核，也不证明Excel正常模式打开无修复提示。'
    if package_failures:
        return {"passed": False, "rows": len(data.get("units", [])), "failures": failures, "limitations": limitations}
    def need(ok, message):
        if not ok:
            failures.append(message)
    try:
        workbook = load_workbook(path, data_only=False)
    except Exception as error:
        return {"passed": False, "rows": len(data.get("units", [])), "failures": failures + ['无法读取工作簿：' + str(error)], "limitations": limitations}
    need(len(workbook.worksheets) == 1, "默认14列交付应只有一张主表")
    sheet = workbook.worksheets[0]
    need(sheet.max_column == 14, "保存文件列数不是14")
    need({str(m) for m in sheet.merged_cells.ranges} == MERGES, "两行表头合并与合同不同")
    target = data.get("target_environment", {})
    target_name = target.get("name", "") if isinstance(target, dict) else ""
    need(sheet["A1"].value == "编号" and sheet["H1"].value == "是否适用于" + target_name, "表头与目标环境不一致")
    language = data.get('language_plan', {})
    language = language if isinstance(language, dict) else {}
    from_english = language.get('local_text_kind') == 'translated_from_english'
    local_language = str(language.get('local_language', ''))
    local_is_english = local_language.lower() in ENGLISH
    try:
        for column, header in language_headers(data).items():
            need(sheet[column + '2'].value == header, column + '表头与实际来源身份不一致')
    except (ValueError, KeyError, TypeError, IndexError, AttributeError) as error:
        need(False, str(error))
    colors = applicability_colors(data)
    reference = load_workbook(template_path).worksheets[0] if template_path else None
    if reference is not None:
        for column in range(1, 15):
            letter = chr(64 + column)
            actual_dimension = sheet.column_dimensions[letter]
            expected_dimension = reference.column_dimensions[letter]
            need(abs(actual_dimension.width - expected_dimension.width) < 0.01, letter + "列宽改变")
            need(actual_dimension.hidden == expected_dimension.hidden, letter + "隐藏设置改变")
    units = data.get("units", [])
    for offset, unit in enumerate(units, 3):
        expected = tuple(unit.get("values", []))
        actual = tuple(sheet.cell(offset, c).value for c in range(1, 15))
        need(actual == expected, str(unit.get("id")) + "导出后单元格错位或改变")
        for col in range(1, 15):
            cell = sheet.cell(offset, col)
            need(cell.data_type != "f", cell.coordinate + "出现非预期公式")
            need(cell.data_type != "e", cell.coordinate + "出现Excel错误")
            need(bool(cell.alignment.wrap_text), cell.coordinate + "未自动换行")
        if unit.get("values", [None] * 8)[7] in colors:
            need(fill_rgb(sheet.cell(offset, 8).fill) == colors[unit["values"][7]][1:].upper(), "H" + str(offset) + "适用性颜色错误")
        if reference is not None:
            source_row = min(offset, max(3, reference.max_row))
            for column in range(1, 15):
                if column != 8:
                    need(copy(sheet.cell(offset, column).fill) == copy(reference.cell(source_row, column).fill), sheet.cell(offset, column).coordinate + "改变了H列以外的配色")
        height = sheet.row_dimensions[offset].height
        need(height is None or height <= 409.5, str(offset) + "行高超过Excel上限")
    rules_found = set()
    for cf in sheet.conditional_formatting:
        covers_body = any(r.min_col == 8 and r.max_col == 8 and r.min_row <= 3 and r.max_row >= len(units) + 2 for r in cf.sqref.ranges)
        if not covers_body:
            continue
        for rule in sheet.conditional_formatting[cf]:
            for value in ["是", "否"]:
                if (rule.type == "cellIs" and rule.operator == "equal" and rule.formula == ['"' + value + '"']
                        and rule.dxf is not None and rule.dxf.fill is not None
                        and fill_rgb(rule.dxf.fill) == colors.get(value, "")[1:].upper()):
                    rules_found.add(value)
    need(rules_found == {"是", "否"}, "H列缺少随是／否变化的专属颜色规则")
    for rr in sheet.iter_rows(min_row=len(units) + 3):
        need(all(c.value is None for c in rr), "表尾残留其他法规内容：" + str(rr[0].row))
    note = sheet["E3"].comment
    need(note is not None and data.get("regulation", {}).get("source_note", "") in note.text, "首条当地语言单元格缺来源备注或备注错位")
    return {"passed": not failures, "rows": len(units), "failures": failures, "limitations": limitations}

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--manifest", required=True)
    parser.add_argument("--report")
    parser.add_argument("--template", help="The actual template used, for style-preservation checks")
    args = parser.parse_args()
    data = json.loads(Path(args.manifest).read_text(encoding="utf-8-sig"))
    report = verify_workbook(data, args.workbook, args.template)
    if args.report:
        Path(args.report).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf8")
    print(json.dumps(report, ensure_ascii=False))
    raise SystemExit(0 if report["passed"] else 1)

if __name__ == "__main__":
    main()
